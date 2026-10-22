#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build Presentation Script for IT3120 Practice 1: Problem Statement -> User Needs
Author: Business Analyst / System Architect Pair
Target: slides/IT3120_ThucHanh1_General_ProblemStatement.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE (Eye-friendly, Academic, Professional)
# ==============================================================================
NAVY_DARK = RGBColor(15, 43, 72)        # #0F2B48 - Deep Academic Navy
TEAL = RGBColor(13, 148, 136)           # #0D9488 - Soothing Teal (Accent)
BLUE_PRIMARY = RGBColor(37, 99, 235)    # #2563EB - Primary Blue (Sinh viên)
AMBER = RGBColor(194, 98, 14)           # #C2620E - Warm Amber (Giảng viên / Giả định)
PURPLE = RGBColor(124, 58, 237)         # #7C3AED - Purple Slate (Tài vụ / Học phí)
GREEN_SAGE = RGBColor(22, 163, 74)      # #16A34A - Sage Green (Outcomes/Evidence)
TEXT_DARK = RGBColor(30, 41, 59)        # #1E293B - Charcoal Dark Text
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B - Slate Muted Text
CARD_BG = RGBColor(248, 250, 252)       # #F8FAFC - Card Background
CARD_BORDER = RGBColor(203, 213, 225)   # #CBD5E1 - Card Border
WHITE = RGBColor(255, 255, 255)
HERO_BG = RGBColor(240, 249, 255)       # Soft Blue-sky Tint for Hero
HIGHLIGHT_GREEN_BG = RGBColor(240, 253, 244) # Soft Green Tint
HIGHLIGHT_GREEN_BORDER = RGBColor(134, 239, 172)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen aspect ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    total_slides = 8

    def add_header(slide, title_text, category_text="THỰC HÀNH 1: PROBLEM STATEMENT → USER NEEDS"):
        # Header background line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = TEAL
        line.line.color.rgb = TEAL

        # Title box
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(9.5), Inches(0.65))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_HEADING
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        # Category box (right)
        catBox = slide.shapes.add_textbox(Inches(8.5), Inches(0.45), Inches(4.033), Inches(0.45))
        ctf = catBox.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = category_text
        cp.alignment = PP_ALIGN.RIGHT
        cp.font.name = FONT_BODY
        cp.font.size = Pt(10.5)
        cp.font.bold = True
        cp.font.color.rgb = TEAL

    def add_footer(slide, current_slide):
        # Footer thin line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.color.rgb = CARD_BORDER

        # Footer text left
        fBoxL = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(5.0), Inches(0.3))
        ftfL = fBoxL.text_frame
        ftfL.margin_left = ftfL.margin_top = ftfL.margin_right = ftfL.margin_bottom = 0
        pL = ftfL.paragraphs[0]
        pL.text = "IT3120 - Phân tích & Thiết kế Hệ thống | SOICT - ĐHBK Hà Nội"
        pL.font.name = FONT_BODY
        pL.font.size = Pt(9.5)
        pL.font.color.rgb = TEXT_MUTED

        # Footer text center
        fBoxC = slide.shapes.add_textbox(Inches(5.0), Inches(7.1), Inches(5.0), Inches(0.3))
        ftfC = fBoxC.text_frame
        ftfC.margin_left = ftfC.margin_top = ftfC.margin_right = ftfC.margin_bottom = 0
        pC = ftfC.paragraphs[0]
        pC.alignment = PP_ALIGN.CENTER
        pC.text = "Đề tài: Hệ thống Quản lý Đào tạo Đại học theo Học chế Tín chỉ"
        pC.font.name = FONT_BODY
        pC.font.size = Pt(9.5)
        pC.font.color.rgb = TEXT_MUTED

        # Footer text right (Slide number)
        fBoxR = slide.shapes.add_textbox(Inches(10.533), Inches(7.1), Inches(2.0), Inches(0.3))
        ftfR = fBoxR.text_frame
        ftfR.margin_left = ftfR.margin_top = ftfR.margin_right = ftfR.margin_bottom = 0
        pR = ftfR.paragraphs[0]
        pR.alignment = PP_ALIGN.RIGHT
        pR.text = f"Slide {current_slide} / {total_slides}"
        pR.font.name = FONT_BODY
        pR.font.size = Pt(9.5)
        pR.font.bold = True
        pR.font.color.rgb = NAVY_DARK

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        if left_stripe_color:
            stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.07), height)
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = left_stripe_color
            stripe.line.fill.background()
        return card

    # ==============================================================================
    # SLIDE 1: TITLE SLIDE
    # ==============================================================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Top banner background shape
    banner = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(3.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY_DARK
    banner.line.fill.background()

    # Teal accent line under banner
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.6), Inches(13.333), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = TEAL
    accent_bar.line.fill.background()

    # Tag on top
    tag_box = slide1.shapes.add_textbox(Inches(1.0), Inches(0.55), Inches(11.333), Inches(0.4))
    ttf = tag_box.text_frame
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
    tp = ttf.paragraphs[0]
    tp.text = "TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG (SOICT) — ĐẠI HỌC BÁCH KHOA HÀ NỘI"
    tp.font.name = FONT_BODY
    tp.font.size = Pt(11)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(180, 205, 230)

    # Title on banner
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(11.333), Inches(1.8))
    titletf = title_box.text_frame
    titletf.word_wrap = True
    titletf.margin_left = titletf.margin_top = titletf.margin_right = titletf.margin_bottom = 0
    p1 = titletf.paragraphs[0]
    p1.text = "THỰC HÀNH 1: PROBLEM STATEMENT → USER NEEDS"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(29)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = titletf.add_paragraph()
    p2.text = "Đề tài: Hệ thống Quản lý Đào tạo Đại học theo Học chế Tín chỉ"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(17.5)
    p2.font.color.rgb = RGBColor(140, 230, 220)
    p2.space_before = Pt(7)

    p3 = titletf.add_paragraph()
    p3.text = "Môn học: Phân tích và Thiết kế Hệ thống (IT3120)"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12.5)
    p3.font.color.rgb = RGBColor(210, 225, 240)
    p3.space_before = Pt(5)

    # Bottom Content Container (Left: Instructor Principle Quote, Right: Metadata Card)
    add_card(slide1, Inches(1.0), Inches(4.1), Inches(7.4), Inches(2.7), bg_color=WHITE, border_color=CARD_BORDER, left_stripe_color=TEAL)
    quote_box = slide1.shapes.add_textbox(Inches(1.3), Inches(4.3), Inches(6.8), Inches(2.3))
    qtf = quote_box.text_frame
    qtf.word_wrap = True
    qtf.margin_left = qtf.margin_top = qtf.margin_right = qtf.margin_bottom = 0
    qp1 = qtf.paragraphs[0]
    qp1.text = "ĐỊNH HƯỚNG CỐT LÕI TỪ GIẢNG VIÊN HƯỚNG DẪN:"
    qp1.font.name = FONT_HEADING
    qp1.font.size = Pt(12)
    qp1.font.bold = True
    qp1.font.color.rgb = TEAL

    qp2 = qtf.add_paragraph()
    qp2.text = "“Chỉ cần phân tích Problem Statement chung nhất bao quát các vấn đề (problems), giữ độ mở ở mức tổng quan để các buổi sau mới đi vào phân tích sâu (Business Model Canvas / Value Proposition Canvas, UX, Use Case...).“"
    qp2.font.name = FONT_BODY
    qp2.font.size = Pt(13)
    qp2.font.italic = True
    qp2.font.color.rgb = NAVY_DARK
    qp2.space_before = Pt(7)

    qp3 = qtf.add_paragraph()
    qp3.text = "• Trọng tâm: Xây dựng 04 User Needs đại diện 4 chủ thể, đào sâu 01 nhu cầu ưu tiên và phát biểu bài toán chung nhất."
    qp3.font.name = FONT_BODY
    qp3.font.size = Pt(11)
    qp3.font.color.rgb = TEXT_MUTED
    qp3.space_before = Pt(8)

    # Metadata Card (Right)
    add_card(slide1, Inches(8.7), Inches(4.1), Inches(3.633), Inches(2.7), bg_color=CARD_BG, border_color=CARD_BORDER)
    meta_box = slide1.shapes.add_textbox(Inches(8.95), Inches(4.3), Inches(3.15), Inches(2.3))
    mtf = meta_box.text_frame
    mtf.word_wrap = True
    mtf.margin_left = mtf.margin_top = mtf.margin_right = mtf.margin_bottom = 0
    mp1 = mtf.paragraphs[0]
    mp1.text = "THÔNG TIN HỌC VỤ & NHÓM"
    mp1.font.name = FONT_HEADING
    mp1.font.size = Pt(12)
    mp1.font.bold = True
    mp1.font.color.rgb = NAVY_DARK

    items = [
        ("Giảng viên:", "Bộ môn Hệ thống Thông tin"),
        ("Đơn vị đào tạo:", "SOICT — ĐHBK Hà Nội"),
        ("Quy mô nhóm:", "04 Thành viên (4 Phân hệ)"),
        ("Đối tượng nghiên cứu:", "Quản lý đào tạo tín chỉ"),
        ("Phương pháp luận:", "Problem Statement → User Needs")
    ]
    for lbl, val in items:
        p = mtf.add_paragraph()
        p.text = f"{lbl} {val}"
        p.font.name = FONT_BODY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(5)

    # ==============================================================================
    # SLIDE 2: BỐI CẢNH CHUNG & CÁC VẤN ĐỀ BỀ MẶT (SITUATION & PROBLEMS)
    # ==============================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Bối cảnh Chung & Các Vấn đề Bề mặt (Situation & Problems)")
    add_footer(slide2, 2)

    # Situation Banner
    add_card(slide2, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.95), bg_color=WHITE, border_color=CARD_BORDER, left_stripe_color=NAVY_DARK)
    sit_box = slide2.shapes.add_textbox(Inches(1.05), Inches(1.22), Inches(11.2), Inches(0.8))
    stf = sit_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
    sp1 = stf.paragraphs[0]
    sp1.text = "BỐI CẢNH ĐÀO TẠO ĐẠI HỌC THEO HỌC CHẾ TÍN CHỈ:"
    sp1.font.name = FONT_HEADING
    sp1.font.size = Pt(11.5)
    sp1.font.bold = True
    sp1.font.color.rgb = NAVY_DARK

    sp2 = stf.add_paragraph()
    sp2.text = "Mô hình đào tạo theo học chế tín chỉ trao quyền tự chủ học tập cho sinh viên. Với quy mô hơn 35.000 người học và hàng trăm học phần mở đồng thời, việc cá nhân hóa kế hoạch học tập tạo ra áp lực tương tác và đồng bộ dữ liệu khổng lồ giữa 4 chủ thể liên quan."
    sp2.font.name = FONT_BODY
    sp2.font.size = Pt(11)
    sp2.font.color.rgb = TEXT_DARK
    sp2.space_before = Pt(3)

    # 4 Cards for 4 Problems (Sinh viên, Giảng viên, PĐT, Tài vụ)
    card_width_4 = Inches(2.78)
    gap_4 = Inches(0.2)
    top_pos_4 = Inches(2.25)
    card_height_4 = Inches(3.35)

    problems_data = [
        (
            "1. SINH VIÊN (REG)",
            BLUE_PRIMARY,
            [
                "Lúng túng khi tự lập kế hoạch học tập dài hạn qua các kỳ.",
                "Cổng nghẽn mạng lúc mở đợt, dễ lỡ môn tiên quyết cốt lõi.",
                "Bị động trong việc theo dõi tích lũy tín chỉ và nghĩa vụ học phí."
            ]
        ),
        (
            "2. GIẢNG VIÊN (GRD)",
            AMBER,
            [
                "Quản lý danh sách lớp đông sinh viên từ nhiều khóa/ngành rời rạc.",
                "Nhập và tính toán điểm thành phần bằng Excel thủ công dễ sai sót.",
                "Áp lực rà soát và nộp bảng điểm an toàn, đúng hạn quy chế."
            ]
        ),
        (
            "3. PHÒNG ĐÀO TẠO (SCH)",
            TEAL,
            [
                "Khối lượng mở lớp biến động liên tục qua các đợt đăng ký.",
                "Quá tải tiếp nhận đơn xin mở lớp, hủy lớp nộp thủ công.",
                "Nguy cơ vỡ kế hoạch điều phối giảng đường và thời khóa biểu."
            ]
        ),
        (
            "4. PHÒNG TÀI VỤ (FIN)",
            PURPLE,
            [
                "Chậm xác định công nợ học phí do sinh viên thay đổi môn học.",
                "Đối soát sao kê ngân hàng thủ công tiêu tốn nhiều nhân lực.",
                "Chậm thu hồi ngân sách đào tạo để tái đầu tư cơ sở vật chất."
            ]
        )
    ]

    for idx, (p_title, p_color, p_bullets) in enumerate(problems_data):
        c_left = Inches(0.8) + idx * (card_width_4 + gap_4)
        add_card(slide2, c_left, top_pos_4, card_width_4, card_height_4, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=p_color)

        c_box = slide2.shapes.add_textbox(c_left + Inches(0.15), top_pos_4 + Inches(0.15), card_width_4 - Inches(0.25), card_height_4 - Inches(0.3))
        ctf = c_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0

        p = ctf.paragraphs[0]
        p.text = p_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = p_color

        for b in p_bullets:
            pb = ctf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(6)

    # Bottom Core Problem Callout Box
    add_card(slide2, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.15), bg_color=HIGHLIGHT_GREEN_BG, border_color=HIGHLIGHT_GREEN_BORDER)
    core_box = slide2.shapes.add_textbox(Inches(1.05), Inches(5.83), Inches(11.2), Inches(0.95))
    cortf = core_box.text_frame
    cortf.word_wrap = True
    cortf.margin_left = cortf.margin_top = cortf.margin_right = cortf.margin_bottom = 0
    cp1 = cortf.paragraphs[0]
    cp1.text = "🎯 ĐIỂM GIAO THOA CỐT LÕI (CORE ROOT PROBLEM):"
    cp1.font.name = FONT_HEADING
    cp1.font.size = Pt(11.5)
    cp1.font.bold = True
    cp1.font.color.rgb = GREEN_SAGE

    cp2 = cortf.add_paragraph()
    cp2.text = "Các vấn đề rời rạc ở trên đều quy tụ về một nguyên nhân gốc: Sự phân tán thông tin và thiếu tính liên thông dữ liệu học vụ theo thời gian thực giữa Sinh viên – Giảng viên – Nhà trường – Tài vụ trong suốt vòng đời đào tạo, gây ma sát quy trình và quá tải vận hành."
    cp2.font.name = FONT_BODY
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = NAVY_DARK
    cp2.space_before = Pt(3)

    # ==============================================================================
    # SLIDE 3: NHIỆM VỤ 1 & 2: 04 USER NEEDS, BỐI CẢNH & TRỞ NGẠI
    # ==============================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Nhiệm vụ 1 & 2: 04 User Needs, Bối cảnh & Trở ngại")
    add_footer(slide3, 3)

    # 4 Horizontal Cards (Rows for 4 Members)
    row_height_4 = Inches(1.22)
    row_gap_4 = Inches(0.12)
    start_top_4 = Inches(1.15)

    rows_data_4 = [
        (
            "1. SINH VIÊN CHÍNH QUY (THÀNH VIÊN A - PHÂN HỆ REG)",
            BLUE_PRIMARY,
            "Chủ động lập kế hoạch học tập, chọn lớp và giữ chỗ kịp thời các học phần cốt lõi theo đúng tiến độ ra trường mà không bị xung đột lịch học.",
            "Cổng mở định kỳ với 35.000 SV truy cập đồng thời; lớp hết chỗ tức thì; hệ thống thiếu cảnh báo tiên quyết và không cho giữ chỗ tạm."
        ),
        (
            "2. GIẢNG VIÊN GIẢNG DẠY (THÀNH VIÊN B - PHÂN HỆ GRD)",
            AMBER,
            "Nắm bắt thông tin lớp giảng dạy, ghi nhận và công bố kết quả đánh giá sinh viên tập trung, tự động hóa tính điểm theo đề cương môn học.",
            "Sĩ số lớp đông từ nhiều khóa/ngành; công cụ quản lý điểm rời rạc; nhập điểm thủ công dễ sai lệch công thức trọng số hoặc mất mạng."
        ),
        (
            "3. CÁN BỘ PHÒNG ĐÀO TẠO (THÀNH VIÊN C - PHÂN HỆ CUR/SCH)",
            TEAL,
            "Quy hoạch kế hoạch mở lớp, điều phối phòng học và kiểm soát tiến độ học vụ toàn trường trên một nguồn dữ liệu thống nhất.",
            "Khối lượng mở lớp biến động lớn; dữ liệu phân tán khiến việc rà soát điều kiện mở lớp và xử lý đơn thư khiếu nại thủ công gây quá tải."
        ),
        (
            "4. CÁN BỘ PHÒNG TÀI VỤ (THÀNH VIÊN D - PHÂN HỆ REG/FIN)",
            PURPLE,
            "Xác định chính xác nghĩa vụ học phí theo số tín chỉ thực và đối soát quyết toán công nợ tự động ngay sau khi đóng cổng đăng ký.",
            "Biến động rút/hủy môn liên tục đầu kỳ; đối soát sao kê thủ công tốn thời gian; chậm dòng tiền thu hồi kinh phí tái đầu tư đào tạo."
        )
    ]

    for idx, (role, color, need_txt, barrier_txt) in enumerate(rows_data_4):
        curr_top = start_top_4 + idx * (row_height_4 + row_gap_4)
        add_card(slide3, Inches(0.8), curr_top, Inches(11.733), row_height_4, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=color)

        c_box = slide3.shapes.add_textbox(Inches(1.05), curr_top + Inches(0.08), Inches(11.2), row_height_4 - Inches(0.16))
        ctf = c_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0

        p = ctf.paragraphs[0]
        p.text = role
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p_need = ctf.add_paragraph()
        p_need.text = f"• Nhu cầu (User Need): {need_txt}"
        p_need.font.name = FONT_BODY
        p_need.font.size = Pt(10.5)
        p_need.font.bold = True
        p_need.font.color.rgb = TEXT_DARK
        p_need.space_before = Pt(2)

        p_bar = ctf.add_paragraph()
        p_bar.text = f"• Bối cảnh & Trở ngại: {barrier_txt}"
        p_bar.font.name = FONT_BODY
        p_bar.font.size = Pt(10)
        p_bar.font.color.rgb = TEXT_MUTED
        p_bar.space_before = Pt(2)

    # Principle tip banner at bottom
    add_card(slide3, Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.5), bg_color=WHITE, border_color=CARD_BORDER)
    tip_box = slide3.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.38))
    ttf = tip_box.text_frame
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
    tp = ttf.paragraphs[0]
    tp.text = "📌 Quy chuẩn SOICT: 04 User Needs tập trung vào ĐỘNG TỪ hành động (mục tiêu con người), phân công tương ứng 4 thành viên và 4 phân hệ lớn."
    tp.font.name = FONT_BODY
    tp.font.size = Pt(10.5)
    tp.font.italic = True
    tp.font.color.rgb = NAVY_DARK

    # ==============================================================================
    # SLIDE 4: NHIỆM VỤ 3: BẰNG CHỨNG HIỆN CÓ VS. GIẢ ĐỊNH CẦN KIỂM CHỨNG
    # ==============================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Nhiệm vụ 3: Bằng chứng Hiện có vs. Giả định Cần kiểm chứng")
    add_footer(slide4, 4)

    # Guideline banner
    add_card(slide4, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.65), bg_color=WHITE, border_color=CARD_BORDER, left_stripe_color=TEAL)
    g_box = slide4.shapes.add_textbox(Inches(1.05), Inches(1.22), Inches(11.2), Inches(0.5))
    gtf = g_box.text_frame
    gtf.margin_left = gtf.margin_top = gtf.margin_right = gtf.margin_bottom = 0
    gp = gtf.paragraphs[0]
    gp.text = "DẪN LUẬN GIẢNG VIÊN (Slide 15): “Hãy sử dụng dữ kiện của tình huống, không tự coi những suy đoán là sự thật. Mục tiêu của bước này là tạo ra một phát biểu đủ cụ thể để phân tích sâu hơn, nhưng vẫn mở để có thể điều chỉnh khi có bằng chứng mới.”"
    gp.font.name = FONT_BODY
    gp.font.size = Pt(10.5)
    gp.font.italic = True
    gp.font.color.rgb = NAVY_DARK

    col_width = Inches(5.72)
    col_gap = Inches(0.29)
    col_top = Inches(1.95)
    col_height = Inches(4.9)

    # Left Column: Bằng chứng hiện có
    add_card(slide4, Inches(0.8), col_top, col_width, col_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=GREEN_SAGE)
    l_box = slide4.shapes.add_textbox(Inches(1.05), col_top + Inches(0.18), col_width - Inches(0.35), col_height - Inches(0.36))
    ltf = l_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_top = ltf.margin_right = ltf.margin_bottom = 0

    p = ltf.paragraphs[0]
    p.text = "📁 BẰNG CHỨNG HIỆN CÓ (EVIDENCE)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_SAGE

    p_sub = ltf.add_paragraph()
    p_sub.text = "Dữ kiện thực tế xác thực từ quy chế đào tạo tín chỉ Bách khoa:"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(3)

    evidences = [
        ("Cơ chế quản lý theo học phần độc lập:", "Mọi hoạt động đào tạo đều gắn với học phần. Mỗi học phần có điều kiện tiên quyết, sĩ số giới hạn và số tín chỉ quy định chặt chẽ."),
        ("Chuỗi quan hệ nhân quả khép kín:", "Kế hoạch mở lớp của PĐT $\\rightarrow$ Đăng ký môn của Sinh viên $\\rightarrow$ Nhập điểm của Giảng viên $\\rightarrow$ Tính học phí & xét học vụ của Tài vụ/PĐT."),
        ("Tính pháp lý bắt buộc của Quy chế:", "Mọi quy trình xét tốt nghiệp, cấp học bổng hay cảnh báo học vụ đều dựa trên dữ liệu tích lũy chính xác và phải tuân thủ Quy chế đào tạo chính thức.")
    ]
    for title, desc in evidences:
        p = ltf.add_paragraph()
        p.text = f"✔ {title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(10)

        pd = ltf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(2)

    # Right Column: Giả định cần kiểm chứng
    add_card(slide4, Inches(0.8) + col_width + col_gap, col_top, col_width, col_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=AMBER)
    r_box = slide4.shapes.add_textbox(Inches(0.8) + col_width + col_gap + Inches(0.25), col_top + Inches(0.18), col_width - Inches(0.35), col_height - Inches(0.36))
    rtf = r_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = 0

    p = rtf.paragraphs[0]
    p.text = "❓ GIẢ ĐỊNH CẦN KIỂM CHỨNG (ASSUMPTIONS)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AMBER

    p_sub2 = rtf.add_paragraph()
    p_sub2.text = "Các giả thuyết nghiên cứu cần tiếp tục khảo sát người dùng thực tế:"
    p_sub2.font.name = FONT_BODY
    p_sub2.font.size = Pt(10.5)
    p_sub2.font.color.rgb = TEXT_MUTED
    p_sub2.space_before = Pt(3)

    assumptions = [
        ("Giả định 1 (Rào cản lớn nhất của SV):", "Khó khăn chính của sinh viên hiện nay nằm ở khâu 'tiếp cận và hiểu đúng quy chế học vụ' hay nằm ở khâu 'thao tác đăng ký và tranh suất khi nghẽn mạng'?"),
        ("Giả định 2 (Điểm nghẽn thời gian của GV/PĐT):", "Giảng viên và Cán bộ đào tạo đang tốn nhiều thời gian nhất ở công đoạn nào trong kỳ học (mở lớp, nhập điểm, hay xử lý ngoại lệ/phúc khảo)?"),
        ("Giả định 3 (Nhu cầu minh bạch học phí):", "Sinh viên có thực sự cần biết số tiền học phí tạm tính ngay trong lúc chọn môn hay chỉ quan tâm khi phòng Tài vụ xuất hóa đơn chốt đợt?")
    ]
    for title, desc in assumptions:
        p = rtf.add_paragraph()
        p.text = f"❓ {title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(10)

        pd = rtf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(2)

    # ==============================================================================
    # SLIDE 5: NHIỆM VỤ 4: CHỌN NHU CẦU ƯU TIÊN ĐỂ LÀM SÂU
    # ==============================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Nhiệm vụ 4: Chọn Nhu cầu Ưu tiên để Làm Sâu")
    add_footer(slide5, 5)

    # Strategic Choice Intro
    intro5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.4))
    i5tf = intro5.text_frame
    i5tf.margin_left = i5tf.margin_top = i5tf.margin_right = i5tf.margin_bottom = 0
    p = i5tf.paragraphs[0]
    p.text = "Quyết định chiến lược của nhóm: Chọn 01 Nhu cầu cốt lõi nhất làm tâm điểm đào sâu xuyên suốt chuỗi TH2 → TH5:"
    p.font.name = FONT_BODY
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    # Hero Banner (The Priority Need: Sinh viên)
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(2.0), bg_color=HERO_BG, border_color=TEAL)
    hero_box = slide5.shapes.add_textbox(Inches(1.2), Inches(1.75), Inches(10.933), Inches(1.7))
    htf = hero_box.text_frame
    htf.word_wrap = True
    htf.margin_left = htf.margin_top = htf.margin_right = htf.margin_bottom = 0

    p_badge = htf.paragraphs[0]
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.text = "★ NHU CẦU ƯU TIÊN ĐƯỢC CHỌN ĐÀO SÂU (PRIORITY USER NEED) ★"
    p_badge.font.name = FONT_HEADING
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = BLUE_PRIMARY

    p_hero = htf.add_paragraph()
    p_hero.alignment = PP_ALIGN.CENTER
    p_hero.text = "“Nhu cầu của Sinh viên: Chủ động lập kế hoạch học tập, chọn lớp và giữ chỗ kịp thời các học phần cốt lõi theo đúng tiến độ ra trường.”"
    p_hero.font.name = FONT_HEADING
    p_hero.font.size = Pt(19)
    p_hero.font.bold = True
    p_hero.font.color.rgb = NAVY_DARK
    p_hero.space_before = Pt(6)

    p_sub_hero = htf.add_paragraph()
    p_sub_hero.alignment = PP_ALIGN.CENTER
    p_sub_hero.text = "(Đại diện phân hệ Đăng ký học phần REG & Lớp học SCH — Trách nhiệm phối hợp: Thành viên A & Cả nhóm)"
    p_sub_hero.font.name = FONT_BODY
    p_sub_hero.font.size = Pt(11)
    p_sub_hero.font.color.rgb = TEXT_MUTED
    p_sub_hero.space_before = Pt(4)

    # 3 Reasons for selecting this overarching need
    reasons_top = Inches(3.8)
    reasons_height = Inches(3.0)
    rcard_width = Inches(3.75)
    gap = Inches(0.24)

    reasons = [
        (
            "1. 'Tử huyệt' của toàn chuỗi",
            BLUE_PRIMARY,
            "Mắt xích kích hoạt mọi hoạt động học vụ tiếp theo:",
            [
                "Nếu SV không đăng ký được môn, GV không có lớp dạy.",
                "PĐT bị vỡ kế hoạch điều phối giảng đường và TKB.",
                "Phòng Tài vụ không có căn cứ xác định nguồn thu học phí."
            ]
        ),
        (
            "2. Nỗi đau gay gắt & rủi ro cao",
            AMBER,
            "Gây hậu quả nặng nề nhất cho người học và trường:",
            [
                "Nguy cơ chậm tốt nghiệp 1 kỳ (môn mở 1 lần/năm).",
                "Hàng ngàn sinh viên dồn về văn phòng PĐT nộp đơn khiếu nại.",
                "Mất niềm tin vào tính công bằng và minh bạch của hệ thống."
            ]
        ),
        (
            "3. Rất khả thi để làm MVP Test",
            GREEN_SAGE,
            "Hoàn hảo để thiết kế thử nghiệm kiểm chứng:",
            [
                "Dễ dàng dựng bản mẫu Figma (Giỏ môn học Wishlist).",
                "Dễ mời 20 sinh viên Bách khoa tham gia Usability Test.",
                "Đo lường KPI khách quan (Task Completion Rate, SUS Score)."
            ]
        )
    ]

    for idx, (rtitle, rcolor, rsub, rbullets) in enumerate(reasons):
        c_left = Inches(0.8) + idx * (rcard_width + gap)
        add_card(slide5, c_left, reasons_top, rcard_width, reasons_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=rcolor)

        r_box = slide5.shapes.add_textbox(c_left + Inches(0.2), reasons_top + Inches(0.18), rcard_width - Inches(0.35), reasons_height - Inches(0.36))
        rtf = r_box.text_frame
        rtf.word_wrap = True
        rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = 0

        p = rtf.paragraphs[0]
        p.text = rtitle
        p.font.name = FONT_HEADING
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = rcolor

        ps = rtf.add_paragraph()
        ps.text = rsub
        ps.font.name = FONT_BODY
        ps.font.size = Pt(10.5)
        ps.font.bold = True
        ps.font.color.rgb = NAVY_DARK
        ps.space_before = Pt(5)

        for b in rbullets:
            pb = rtf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(4)

    # ==============================================================================
    # SLIDE 6: NHIỆM VỤ 5: PHÁT BIỂU BÀI TOÁN CHUNG NHẤT (PROBLEM STATEMENT)
    # ==============================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Nhiệm vụ 5: Phát biểu Bài toán Chung nhất (Problem Statement)")
    add_footer(slide6, 6)

    # Standard Formula Box
    add_card(slide6, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.85), bg_color=WHITE, border_color=CARD_BORDER, left_stripe_color=NAVY_DARK)
    form_box = slide6.shapes.add_textbox(Inches(1.05), Inches(1.22), Inches(11.2), Inches(0.7))
    ftf = form_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    fp1 = ftf.paragraphs[0]
    fp1.text = "CÔNG THỨC PHÁT BIỂU CHUẨN MỰC (SLIDE 11):"
    fp1.font.name = FONT_HEADING
    fp1.font.size = Pt(11)
    fp1.font.bold = True
    fp1.font.color.rgb = NAVY_DARK

    fp2 = ftf.add_paragraph()
    fp2.text = "[Trong bối cảnh nào] + [Nhóm người dùng nào] + [Cần nhu cầu cốt lõi gì] + [Gặp trở ngại gốc rễ nào] + [Hướng tới mục tiêu giá trị gì]"
    fp2.font.name = FONT_BODY
    fp2.font.size = Pt(11.5)
    fp2.font.bold = True
    fp2.font.color.rgb = TEAL
    fp2.space_before = Pt(3)

    # Main Hero Problem Statement Container
    add_card(slide6, Inches(0.8), Inches(2.15), Inches(11.733), Inches(3.6), bg_color=HERO_BG, border_color=TEAL)
    hero_st_box = slide6.shapes.add_textbox(Inches(1.15), Inches(2.35), Inches(11.033), Inches(3.2))
    hstf = hero_st_box.text_frame
    hstf.word_wrap = True
    hstf.margin_left = hstf.margin_top = hstf.margin_right = hstf.margin_bottom = 0

    hp_title = hstf.paragraphs[0]
    hp_title.text = "BẢN PHÁT BIỂU VẤN ĐỀ NGƯỜI DÙNG TỔNG QUÁT (GENERAL USER PROBLEM STATEMENT):"
    hp_title.font.name = FONT_HEADING
    hp_title.font.size = Pt(12.5)
    hp_title.font.bold = True
    hp_title.font.color.rgb = TEAL

    statement_parts = (
        "“Trong môi trường giáo dục đại học vận hành theo học chế tín chỉ (Bối cảnh), "
        "sinh viên chính quy cùng với giảng viên và cán bộ quản lý (Đối tượng người dùng) "
        "cần một phương thức quản lý kế hoạch học tập và đăng ký giữ chỗ học phần xuyên suốt, minh bạch và tức thời (Nhu cầu cốt lõi), "
        "bởi vì đặc thù số lượng sinh viên truy cập cao điểm lớn trong khi thông tin lớp học và điều kiện ràng buộc bị phân tán, gây nghẽn tắc hệ thống và làm quá tải việc xử lý đơn thư thủ công (Trở ngại cốt lõi), "
        "nhằm bảo đảm sinh viên tích lũy đủ học phần đúng hạn ra trường và nâng cao hiệu quả vận hành của Nhà trường (Mục tiêu giá trị).”"
    )
    hp_body = hstf.add_paragraph()
    hp_body.text = statement_parts
    hp_body.font.name = FONT_BODY
    hp_body.font.size = Pt(15)
    hp_body.font.color.rgb = NAVY_DARK
    hp_body.space_before = Pt(9)

    # Methodology Note at bottom
    add_card(slide6, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.9), bg_color=HIGHLIGHT_GREEN_BG, border_color=HIGHLIGHT_GREEN_BORDER)
    note_box = slide6.shapes.add_textbox(Inches(1.05), Inches(6.05), Inches(11.2), Inches(0.7))
    ntf = note_box.text_frame
    ntf.word_wrap = True
    ntf.margin_left = ntf.margin_top = ntf.margin_right = ntf.margin_bottom = 0
    np1 = ntf.paragraphs[0]
    np1.text = "💡 NHẬN ĐỊNH PHƯƠNG PHÁP LUẬN:"
    np1.font.name = FONT_HEADING
    np1.font.size = Pt(11)
    np1.font.bold = True
    np1.font.color.rgb = GREEN_SAGE

    np2 = ntf.add_paragraph()
    np2.text = "Phát biểu trên vừa đủ cụ thể khi neo chặt vào điểm nghẽn đăng ký môn của Sinh viên, vừa đủ khái quát để duy trì tính liên thông với công tác quản trị của Giảng viên, Nhà trường và Tài vụ mà không bị 'thiên kiến giải pháp' (Solution Bias)."
    np2.font.name = FONT_BODY
    np2.font.size = Pt(10.5)
    np2.font.color.rgb = TEXT_DARK
    np2.space_before = Pt(2)

    # ==============================================================================
    # SLIDE 7: PHÂN CÔNG TRÁCH NHIỆM SONG SONG THEO MÔ HÌNH PHỄU HỘI TỤ
    # ==============================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Phân công Trách nhiệm Song song theo Mô hình Phễu Hội tụ")
    add_footer(slide7, 7)

    # Strategy intro banner
    add_card(slide7, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.65), bg_color=WHITE, border_color=CARD_BORDER, left_stripe_color=TEAL)
    m7_box = slide7.shapes.add_textbox(Inches(1.05), Inches(1.22), Inches(11.2), Inches(0.5))
    m7tf = m7_box.text_frame
    m7tf.margin_left = m7tf.margin_top = m7tf.margin_right = m7tf.margin_bottom = 0
    mp7 = m7tf.paragraphs[0]
    mp7.text = "MÔ HÌNH PHỄU HỘI TỤ (FUNNEL MODEL): Bảo đảm 4 thành viên làm việc song song 100% (Zero Waiting), không chồng chéo và đúng chuẩn 01 MVP của đề bài."
    mp7.font.name = FONT_BODY
    mp7.font.size = Pt(10.5)
    mp7.font.bold = True
    mp7.font.color.rgb = NAVY_DARK

    col7_width = Inches(5.72)
    col7_gap = Inches(0.29)
    col7_top = Inches(1.95)
    col7_height = Inches(4.9)

    # Left Column: Pha 1 - Phân kỳ song song theo 4 góc nhìn
    add_card(slide7, Inches(0.8), col7_top, col7_width, col7_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=BLUE_PRIMARY)
    p1_box = slide7.shapes.add_textbox(Inches(1.05), col7_top + Inches(0.18), col7_width - Inches(0.35), col7_height - Inches(0.36))
    p1tf = p1_box.text_frame
    p1tf.word_wrap = True
    p1tf.margin_left = p1tf.margin_top = p1tf.margin_right = p1tf.margin_bottom = 0

    p = p1tf.paragraphs[0]
    p.text = "PHA 1: PHÂN KỲ SONG SONG (TH1 → TH3)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    ps1 = p1tf.add_paragraph()
    ps1.text = "4 bạn mở chung 1 file Docs/Slide, cùng điền dữ liệu theo 4 vai trò:"
    ps1.font.name = FONT_BODY
    ps1.font.size = Pt(10.5)
    ps1.font.color.rgb = TEXT_MUTED
    ps1.space_before = Pt(3)

    pha1_items = [
        ("Thành viên A (Lát cắt Sinh viên - REG):", "Soạn User Need 1, nỗi đau nghẽn mạng, giải pháp Giỏ môn học và các rào cản đăng ký môn."),
        ("Thành viên B (Lát cắt Giảng viên - GRD):", "Soạn User Need 2, nỗi đau điểm số phân tán, sai lệch công thức trọng số và rủi ro chậm nộp sổ điểm."),
        ("Thành viên C (Lát cắt Phòng Đào tạo - SCH):", "Soạn User Need 3, nỗi đau vỡ kế hoạch lớp học, quá tải xử lý đơn từ khiếu nại mở thêm lớp."),
        ("Thành viên D (Lát cắt Phòng Tài vụ - FIN):", "Soạn User Need 4, nỗi đau chậm tiến độ công nợ học phí và quá tải đối soát sao kê ngân hàng.")
    ]
    for role_title, role_desc in pha1_items:
        p = p1tf.add_paragraph()
        p.text = role_title
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(8)

        pd = p1tf.add_paragraph()
        pd.text = role_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(1)

    # Right Column: Pha 2 - Hội tụ chuyên môn cho 01 MVP duy nhất
    add_card(slide7, Inches(0.8) + col7_width + col7_gap, col7_top, col7_width, col7_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=GREEN_SAGE)
    p2_box = slide7.shapes.add_textbox(Inches(0.8) + col7_width + col7_gap + Inches(0.25), col7_top + Inches(0.18), col7_width - Inches(0.35), col7_height - Inches(0.36))
    p2tf = p2_box.text_frame
    p2tf.word_wrap = True
    p2tf.margin_left = p2tf.margin_top = p2tf.margin_right = p2tf.margin_bottom = 0

    p = p2tf.paragraphs[0]
    p.text = "PHA 2: HỘI TỤ CHUYÊN MÔN CHO 01 MVP (TH4 → TH5)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_SAGE

    ps2 = p2tf.add_paragraph()
    ps2.text = "Cùng xúm vào làm sâu DUY NHẤT 1 MVP Đăng ký học phần (Wishlist):"
    ps2.font.name = FONT_BODY
    ps2.font.size = Pt(10.5)
    ps2.font.color.rgb = TEXT_MUTED
    ps2.space_before = Pt(3)

    pha2_items = [
        ("Thành viên A (UX/UI Designer):", "Thiết kế 3 khung màn hình Figma/wireframe: (1) Thêm môn vào giỏ, (2) Xem mô phỏng TKB, (3) Xác nhận 1 chạm."),
        ("Thành viên B (User Researcher):", "Soạn kịch bản Usability Test cho 20 sinh viên và thiết kế 5 câu hỏi phỏng vấn cảm xúc/độ tự tin."),
        ("Thành viên C (Data/QA Analyst):", "Xây dựng bảng chỉ số đo lường KPI: Task Completion Rate $\\ge 90\\%$, thời gian $\\le 2$ phút, điểm hài lòng SUS $\\ge 80$."),
        ("Thành viên D (Lead Documenter):", "Đóng gói toàn bộ tài liệu từ TH1 đến TH5 thành slide PowerPoint và hoàn thiện hồ sơ System Request.")
    ]
    for role_title, role_desc in pha2_items:
        p = p2tf.add_paragraph()
        p.text = role_title
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(8)

        pd = p2tf.add_paragraph()
        pd.text = role_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(1)

    # ==============================================================================
    # SLIDE 8: TỔNG KẾT ĐẦU RA BÀI TẬP 1 & LỘ TRÌNH TIẾP THEO
    # ==============================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. Tổng kết Đầu ra Bài tập 1 & Lộ trình Tiếp theo")
    add_footer(slide8, 8)

    col8_width = Inches(5.72)
    col8_top = Inches(1.25)
    col8_height = Inches(4.8)

    # Left Column: Deliverables
    add_card(slide8, Inches(0.8), col8_top, col8_width, col8_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=GREEN_SAGE)
    dl_box = slide8.shapes.add_textbox(Inches(1.05), col8_top + Inches(0.2), col8_width - Inches(0.35), col8_height - Inches(0.4))
    dltf = dl_box.text_frame
    dltf.word_wrap = True
    dltf.margin_left = dltf.margin_top = dltf.margin_right = dltf.margin_bottom = 0

    p = dltf.paragraphs[0]
    p.text = "✅ 03 ĐẦU RA ĐẠT ĐƯỢC (DELIVERABLES)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_SAGE

    deliverables = [
        ("1. Danh sách 04 User Needs cốt lõi:", "Phản ánh đúng 4 góc nhìn hành vi khách quan (Sinh viên, Giảng viên, Cán bộ đào tạo, Cán bộ tài vụ) khớp 1-1 với 4 thành viên."),
        ("2. Một Nhu cầu Ưu tiên & Problem Statement:", "Đào sâu bài toán sinh tử về đăng ký môn học và giữ chỗ học phần cốt lõi của sinh viên."),
        ("3. Danh mục 03 Giả định & Ma trận Phân công:", "Xác định rõ các câu hỏi mở cần điều tra và quy trình làm việc song song không nghẽn cổ chai.")
    ]
    for d_title, d_desc in deliverables:
        p = dltf.add_paragraph()
        p.text = d_title
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(14)

        pd = dltf.add_paragraph()
        pd.text = d_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(3)

    # Right Column: Next Steps Roadmap
    add_card(slide8, Inches(0.8) + col8_width + col_gap, col8_top, col8_width, col8_height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe_color=BLUE_PRIMARY)
    nx_box = slide8.shapes.add_textbox(Inches(0.8) + col8_width + col_gap + Inches(0.25), col8_top + Inches(0.2), col8_width - Inches(0.35), col8_height - Inches(0.4))
    nxtf = nx_box.text_frame
    nxtf.word_wrap = True
    nxtf.margin_left = nxtf.margin_top = nxtf.margin_right = nxtf.margin_bottom = 0

    p = nxtf.paragraphs[0]
    p.text = "🚀 LỘ TRÌNH PHÂN TÍCH SÂU TIẾP THEO"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    steps = [
        ("Thực hành 2: User Problem Statement Canvas", "Hoàn thiện khung Canvas 6 ô bóc tách bối cảnh, nguyên nhân gốc rễ và giải pháp thay thế cho nhánh Sinh viên."),
        ("Thực hành 3: Xác định Business Problem", "Thiết lập chuỗi nhân quả: User Problem $\\rightarrow$ Operational Impact $\\rightarrow$ Business Problem $\\rightarrow$ Desired Outcomes."),
        ("Thực hành 4 & 5: Lean UX Canvas & MVP Plan", "Tổng hợp 8 ô Canvas, viết giả thuyết và thiết kế kịch bản thử nghiệm Usability Test cho 20 sinh viên trên Figma.")
    ]
    for s_title, s_desc in steps:
        p = nxtf.add_paragraph()
        p.text = s_title
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_before = Pt(14)

        pd = nxtf.add_paragraph()
        pd.text = s_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(3)

    # Thank you footer card
    add_card(slide8, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.65), bg_color=WHITE, border_color=CARD_BORDER)
    thx_box = slide8.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.3), Inches(0.45))
    ttf = thx_box.text_frame
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
    tp = ttf.paragraphs[0]
    tp.alignment = PP_ALIGN.CENTER
    tp.text = "Nhóm xin chân thành cảm ơn Thầy/Cô và các bạn đã chú ý lắng nghe! Sẵn sàng nhận câu hỏi thảo luận."
    tp.font.name = FONT_BODY
    tp.font.size = Pt(11.5)
    tp.font.bold = True
    tp.font.color.rgb = NAVY_DARK

    # ==============================================================================
    # SAVE PRESENTATION
    # ==============================================================================
    out_dir = os.path.join(os.path.dirname(__file__), "..", "slides")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "IT3120_ThucHanh1_General_ProblemStatement.pptx")
    prs.save(out_path)
    print(f"[SUCCESS] PowerPoint created successfully at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    create_presentation()
